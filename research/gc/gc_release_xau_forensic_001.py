#!/usr/bin/env python3
from __future__ import annotations

import io
import json
import re
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

OWNER='chepigga'; REPO='ResearchOS'; TAG='GC'
API=f'https://api.github.com/repos/{OWNER}/{REPO}/releases/tags/{TAG}'
OUT=Path('research/gc')


def get_json(url):
    req=urllib.request.Request(url,headers={'User-Agent':'ResearchOS-GC-Forensic/1.0'})
    with urllib.request.urlopen(req,timeout=60) as r:
        return json.load(r)


def get_bytes(url):
    req=urllib.request.Request(url,headers={'User-Agent':'ResearchOS-GC-Forensic/1.0'})
    with urllib.request.urlopen(req,timeout=120) as r:
        return r.read()


def main():
    rel=get_json(API)
    assets={a['name']:a for a in rel['assets']}
    xname='ReportTester-1514598162_xau.xlsx'
    if xname not in assets:
        raise RuntimeError('XAU report asset missing')
    raw=get_bytes(assets[xname]['browser_download_url'])
    wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    dump=[]
    trade_rows=[]
    for ws in wb.worksheets:
        dump.append(f'### SHEET {ws.title} dim={ws.max_row}x{ws.max_column}')
        for ridx,row in enumerate(ws.iter_rows(values_only=True),start=1):
            vals=['' if v is None else str(v) for v in row]
            non=[v for v in vals if v!='']
            if not non: continue
            line='\t'.join(vals)
            dump.append(f'{ridx}\t{line}')
            low=' '.join(non).lower()
            if ('buy' in low or 'sell' in low) and re.search(r'2026[.\-/]', low):
                trade_rows.append({'sheet':ws.title,'row':ridx,'values':non})
    (OUT/'GC_RELEASE_XAU_REPORT_FLAT.txt').write_text('\n'.join(dump),encoding='utf-8')

    mq_names=[n for n in assets if n.lower().endswith('.mq5') and ('Rithmic' in n or 'XAU_Context' in n)]
    scans={}
    pats=['AEIF','q10','q90','q75','sell_loc','buy_loc','lower20','upper20','delta_frac','ATR20','price impact','impact','confirmation','aggressor']
    for name in mq_names:
        txt=get_bytes(assets[name]['browser_download_url']).decode('utf-8','replace')
        hits=[]
        for no,line in enumerate(txt.splitlines(),1):
            if any(p.lower() in line.lower() for p in pats):
                hits.append({'line':no,'text':line[:500]})
        scans[name]={'bytes':len(txt.encode('utf-8')),'hits':hits[:500]}
    result={
        'release':rel['html_url'],
        'xlsx_asset':xname,
        'sheets':wb.sheetnames,
        'trade_like_rows':trade_rows,
        'mq5_scans':scans,
    }
    (OUT/'GC_RELEASE_XAU_FORENSIC_001.json').write_text(json.dumps(result,indent=2,ensure_ascii=False,default=str),encoding='utf-8')
    md=['# GC_RELEASE_XAU_FORENSIC_001','',f"- Workbook sheets: {', '.join(wb.sheetnames)}",f'- Trade-like rows: {len(trade_rows)}','']
    md+=['## Trade-like rows','']+[f"- `{r['sheet']}:{r['row']}` — {' | '.join(r['values'])}" for r in trade_rows[:200]]
    md+=['','## MQL5 scans','']
    for n,s in scans.items():
        md.append(f'### {n}')
        md.append(f"Hits: {len(s['hits'])}")
        for h in s['hits'][:100]: md.append(f"- L{h['line']}: `{h['text'].replace('`','')}`")
        md.append('')
    (OUT/'GC_RELEASE_XAU_FORENSIC_001.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    print('\n'.join(md))

if __name__=='__main__': main()
